"""Extrae transcripción suministrada por coordinación. No valida equivalencia con 2026."""
from pathlib import Path
import argparse,json,hashlib,re,ast
from openpyxl import load_workbook
p=argparse.ArgumentParser(description=__doc__);p.add_argument('auditoria',type=Path);p.add_argument('fuente_backend',type=Path);p.add_argument('destino',type=Path);a=p.parse_args()
w=load_workbook(a.auditoria,data_only=True);backend=a.fuente_backend.read_text(encoding='utf-8-sig')
def old(name):return ast.literal_eval(re.search(r'const '+name+r'=(\[.*?\]);',backend,re.S)[1])
previous_g=old('PPS_EVAL_GENERICOS_2026');previous_h=old('PPS_EVAL_DESARROLLO_HUMANO_2026')
g=list(w['Genéricas'].iter_rows(min_row=2,values_only=True));h=list(w['Desarrollo humano'].iter_rows(min_row=2,values_only=True));assert len(g)==16 and len(h)==15
assert [r[0] for r in g]==list(range(1,17)) and [r[0] for r in h]==list(range(1,16))
control=dict((r[0],r[1]) for r in w['Control documental'].iter_rows(min_row=2,values_only=True) if r[0])
doc={'instrumento':'IEPP-01-EC','versionFuente':2025,'contrasteDeclarado':2023,'equivalencia2026Confirmada':False,'origen':'Transcripción aportada por coordinación; no se inspeccionó directamente el archivo original en Drive en esta sesión.','archivoAuditoria':a.auditoria.name,'sha256Auditoria':hashlib.sha256(a.auditoria.read_bytes()).hexdigest(),'urlFuente':control['Documento fuente'],'referenciasFilasAproximadas':True,
'genericas':[{'numero':r[0],'tituloAnterior':previous_g[i],'indicador':r[1],'referencia':r[3]} for i,r in enumerate(g)],
'desarrolloHumano':[{'numero':r[0],'tituloAnterior':previous_h[i],'encabezado':r[1],'descripcion':r[2],'referencia':r[3]} for i,r in enumerate(h)]}
assert all(isinstance(x['indicador'],str) and x['indicador'] for x in doc['genericas'])
assert all(isinstance(x['encabezado'],str) and isinstance(x['descripcion'],str) and x['descripcion'] for x in doc['desarrolloHumano'])
a.destino.mkdir(parents=True,exist_ok=True);text=json.dumps(doc,ensure_ascii=False,indent=2)+'\n';(a.destino/'textos-documentados-2025.json').write_text(text)
(a.destino/'textos-documentados-2025.js').write_text('/* Fuente documental 2025; equivalencia con 2026 no confirmada. */\nwindow.PPS_EVAL_TEXTOS_DOCUMENTADOS='+text.rstrip()+';\n')
print('Extraídos 16 indicadores y 15 encabezados/descripciones; procedencia 2025 conservada.')
