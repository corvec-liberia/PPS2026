"""Contrasta la presentación contra el Excel aportado, sin afirmar equivalencia 2026."""
from pathlib import Path
import sys,json,subprocess
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from guardado_evaluaciones_test import MOCK
root=Path(__file__).resolve().parents[1];source=Path(sys.argv[1]).read_text().replace('<?= PPS_V11_PUBLICO ?>','');fixture=Path(sys.argv[2]);audit=load_workbook(sys.argv[3],data_only=True)
generic=list(audit['Genéricas'].iter_rows(min_row=2,values_only=True));human=list(audit['Desarrollo humano'].iter_rows(min_row=2,values_only=True))
module=(fixture/'Modulo_Docentes_PPS_2026.js').read_text(encoding='utf-8-sig')+'\n'+(fixture/'Untitled.js').read_text(encoding='utf-8-sig')
expr='JSON.stringify(Object.fromEntries(Object.keys(PPS_EVAL_CATALOGO).map(k=>[k,{criterios:ppsEvalCriterios_(k),instrumento:PPS_EVAL_CATALOGO[k].instrumento}])))'
runner="const vm=require('node:vm');let c={};vm.createContext(c);vm.runInContext("+json.dumps(module)+",c);process.stdout.write(vm.runInContext("+json.dumps(expr)+",c));"
catalogs=json.loads(subprocess.check_output(['node','-e',runner],text=True))
mock=MOCK[MOCK.index('window.calls='):MOCK.index('window.openFixture=')].replace('},10);','},window.mockLatency||10);')
html=source.replace('<head>','<head><script>'+mock+'</script>',1)
with sync_playwright() as p:
 browser=p.chromium.launch(executable_path='/usr/bin/chromium',args=['--no-sandbox'])
 page=browser.new_page(viewport={'width':360,'height':800},has_touch=True);page.route('**/*',lambda r:r.abort());errors=[];page.on('pageerror',lambda e:errors.append(str(e)));page.set_content(html)
 def open_case(code='ECSC',number=1,answers=None):
  page.evaluate('''([code,n,answers,data])=>{document.getElementById('loginScreen').style.display='none';PPS_EVAL_ACTUAL={numero:n,idAsignacion:'ASG-SYNTHETIC',codigo:code,estudiante:'Revisión ficticia',empresa:'Empresa ficticia',instrumento:data.instrumento,criterios:data.criterios,respuestas:answers||{},estado:'Borrador',puedeGuardar:true,demo:false};ppsEvalDibujar(PPS_EVAL_ACTUAL);}''',[code,number,answers or {},catalogs[code]])
 covered=0
 for code,data in catalogs.items():
  open_case(code);criteria=data['criterios'];technical=len(criteria)-31;scope=data['instrumento']=='IEPP-01-EC'
  for i in range(len(criteria)):
   page.evaluate('i=>ppsEvalPanel().ppsStepGoTo(i)',i)
   title=page.locator('.pps-step-card h3').text_content();description=page.locator('.pps-step-description').text_content()
   if i<technical:assert title==criteria[i] and description==''
   elif scope and i<technical+16:assert title==generic[i-technical][1] and description==''
   elif scope:assert title==human[i-technical-16][1] and description==human[i-technical-16][2]
   else:assert title==criteria[i].split(' - ',1)[1] and description==''
  if scope:covered+=1;assert '2025' in page.locator('.pps-document-source').inner_text();assert 'pendiente' in page.locator('.pps-document-source').inner_text()
  else:assert page.locator('.pps-document-source').count()==0
 print('PASS: 31 textos literales cotejados con Excel en',covered,'instrumentos EC; IEPP-01 y criterios técnicos intactos')
 open_case();technical=len(catalogs['ECSC']['criterios'])-31
 assert page.evaluate('PPS_EVAL_TEXTOS_DOCUMENTADOS.equivalencia2026Confirmada') is False
 for n in [1,2]:
  open_case(number=n,answers={str(technical):'0',str(technical+16):'NA'})
  page.evaluate('i=>ppsEvalPanel().ppsStepGoTo(i)',technical);assert page.locator('.pps-step-number').inner_text()=='Criterio 1 de 16';assert page.locator('input[value="0"]').is_checked()
  page.evaluate('i=>ppsEvalPanel().ppsStepGoTo(i)',technical+16);assert page.locator('.pps-step-number').inner_text()=='Criterio 1 de 15';assert page.locator('input[value="NA"]').is_checked()
  page.evaluate('ppsEvalGuardar(false)');page.wait_for_function('n=>records["ASG-SYNTHETIC:"+n]?.estado==="Borrador"',arg=n)
  assert page.evaluate('n=>records["ASG-SYNTHETIC:"+n].respuestas',n)=={str(technical):'0',str(technical+16):'NA'}
 print('PASS: numeración por bloque, reapertura de 0/N/A y claves originales en evaluaciones 1 y 2')
 open_case(number=1,answers={str(technical):'0',str(technical+16):'NA'})
 printed=page.evaluate('answers=>ppsEvalContenidoImpresion(PPS_EVAL_ACTUAL,answers)',{str(technical):'0',str(technical+16):'NA'})
 printpage=browser.new_page();printpage.set_content(printed)
 assert printpage.locator('tbody tr.block').count()==3
 blocks=printpage.locator('tbody tr:not(.block)').evaluate_all('rows=>rows.map(r=>r.children[0].textContent)')
 assert blocks==list(map(str,range(1,technical+1)))+list(map(str,range(1,17)))+list(map(str,range(1,16)))
 assert 'Equivalencia con 2026 pendiente' in printpage.locator('body').inner_text()
 for row in generic:assert row[1] in printpage.locator('body').text_content()
 for row in human:assert row[1] in printpage.locator('body').text_content() and row[2] in printpage.locator('body').text_content()
 assert printpage.locator('td.score').filter(has_text='X').count()==2
 out=root.parent/'reports'/'criterios-documentados';out.mkdir(parents=True,exist_ok=True);printpage.pdf(path=str(out/'Revision_bloques_IEPP_01_EC_2025.pdf'),format='A4');printpage.close()
 print('PASS: impresión reinicia los números en cada bloque, incluye todos los textos y conserva 0/N/A')
 # Cierre seguro ante una correspondencia distinta: no completar textos por posición sin verificar título.
 changed=list(catalogs['ECSC']['criterios']);changed[technical]='COMPETENCIAS GENERICAS - Criterio diferente'
 result=page.evaluate('([c,i])=>ppsEvalReferenciaBloque(c,i,"IEPP-01-EC")',[changed,technical]);assert not result['documentado'];assert result['texto']=='Criterio diferente'
 print('PASS: correspondencia por instrumento, orden y título; no se inventa contenido ante una discrepancia')
 for width,height in [(320,640),(360,800),(768,1024)]:
  page.set_viewport_size({'width':width,'height':height});open_case(answers={str(i):'2' for i in range(len(catalogs['ECSC']['criterios']))});page.evaluate('ppsEvalPanel().ppsStepGoTo(24);mockLatency=10000;ppsEvalGuardar(false)')
  r=page.locator('[data-n]').bounding_box();x,y=r['x']+r['width']/2,r['y']+r['height']/2
  page.touchscreen.tap(x,y);assert page.locator('.pps-step-number').inner_text()==f'Criterio {26-technical} de 16'
  new=page.locator('[data-n]').bounding_box();assert abs(new['y']-r['y'])<=2
  assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
  # Mantener cada prueba en su propia evaluación después de que confirme la petición simulada.
  page.wait_for_function('!ppsV11Busy',timeout=15000)
 print('PASS: textos largos y guardado lento conservan botones estables en teléfono y tablet')
 assert not errors,errors;page.close();browser.close()
